import docx
import re

doc = docx.Document('storage/app/templates/Form_7.2-1_Permohonan.docx')

# 1. Delete the first table if it contains KEMENTERIAN
if len(doc.tables) > 0:
    tbl = doc.tables[0]
    text = ""
    for r in tbl.rows:
        for c in r.cells:
            text += c.text + " "
    if 'KEMENTERIAN' in text or 'SDLP' in text:
        tbl._element.getparent().remove(tbl._element)
        print("Removed KEMENTERIAN table.")

# 2. Remove KOP PERUSAHAAN box (might be a table too)
tables_to_remove = []
for tbl in doc.tables:
    text = ""
    for r in tbl.rows:
        for c in r.cells:
            text += c.text + " "
    if 'KOP PERUSAHAAN' in text:
        tables_to_remove.append(tbl)

for tbl in tables_to_remove:
    tbl._element.getparent().remove(tbl._element)
    print("Removed KOP PERUSAHAAN table.")

# 3. Fix split placeholders by consolidating text into first run
for p in doc.paragraphs:
    if 'KOP PERUSAHAAN' in p.text:
        p.text = p.text.replace('KOP PERUSAHAAN', '')
        
    if '${' in p.text and '}' in p.text:
        full_text = p.text
        if len(p.runs) > 0:
            p.runs[0].text = full_text
            for i in range(1, len(p.runs)):
                p.runs[i].text = ''

# Also do this for cells in tables
for tbl in doc.tables:
    for row in tbl.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                if 'KOP PERUSAHAAN' in p.text:
                    p.text = p.text.replace('KOP PERUSAHAAN', '')
                if '${' in p.text and '}' in p.text:
                    full_text = p.text
                    if len(p.runs) > 0:
                        p.runs[0].text = full_text
                        for i in range(1, len(p.runs)):
                            p.runs[i].text = ''

doc.save('storage/app/templates/Form_7.2-1_Permohonan.docx')
print("Cleaned template!")
